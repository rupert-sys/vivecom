"""
F1-17 / HU-A11: exportar estado de cuenta, gastos y presupuesto vs. real a
Excel. Se valida abriendo el .xlsx devuelto con openpyxl (round-trip) en vez
de solo revisar los bytes crudos.
"""

import io
from contextlib import asynccontextmanager
from datetime import datetime
from unittest.mock import patch

import openpyxl
import pytest

from app.services import deposit_processing_service as dps
from app.services.payment_providers.base import DepositoRecibido


@pytest.fixture()
def deposit_env(client):
    @asynccontextmanager
    async def fake_tenant_session(schema_name):
        async with client.db_session_factory() as session:
            yield session

    @asynccontextmanager
    async def fake_control_session():
        async with client.db_session_factory() as session:
            yield session

    with patch.object(dps, "tenant_session", fake_tenant_session):
        yield fake_control_session


def _deposito(referencia="0000001", clave="STP-001", monto=1500.00):
    return DepositoRecibido(
        monto=monto,
        referencia_numerica=referencia,
        clave_rastreo=clave,
        fecha=datetime(2026, 9, 15, 10, 0, 0),
        cuenta_beneficiaria="012180001547896321",
    )


def _leer_xlsx(response) -> openpyxl.Workbook:
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return openpyxl.load_workbook(io.BytesIO(response.content))


@pytest.mark.asyncio
async def test_export_account_statements_includes_charges_and_payments(client, deposit_env):
    prop = client.post("/properties", json={"identificador": "Casa 1"}).json()
    client.post("/fees", json={"monto": 1500.00, "periodicidad": "mensual", "activa_desde": "2026-09-01"})
    client.post("/fees/generate-charges", params={"periodo": "2026-09-01"})

    async with deposit_env() as control_db:
        await dps.process_incoming_deposit(_deposito(), control_db)

    wb = _leer_xlsx(client.get("/reports/account-statements/export"))
    assert wb.sheetnames == ["Cargos", "Pagos"]

    cargos = wb["Cargos"]
    assert cargos.cell(row=1, column=1).value == "Vivienda"
    assert cargos.cell(row=2, column=1).value == "Casa 1"
    assert cargos.cell(row=2, column=6).value == "pagado"

    pagos = wb["Pagos"]
    assert pagos.cell(row=2, column=1).value == "Casa 1"
    assert pagos.cell(row=2, column=3).value == 1500.00

    filtrado = _leer_xlsx(client.get("/reports/account-statements/export", params={"property_id": prop["id"]}))
    assert filtrado["Cargos"].cell(row=2, column=1).value == "Casa 1"


def test_export_expenses_respects_filters(client):
    client.post(
        "/expenses",
        json={
            "categoria": "Mantenimiento", "monto": 1200.00,
            "comprobante_url": "https://x/a.pdf", "fecha": "2026-09-05",
        },
    )
    client.post(
        "/expenses",
        json={
            "categoria": "Seguridad", "monto": 800.00,
            "comprobante_url": "https://x/b.pdf", "fecha": "2026-08-01",
        },
    )

    wb = _leer_xlsx(client.get("/reports/expenses/export", params={"categoria": "Mantenimiento"}))
    hoja = wb["Gastos"]
    assert hoja.cell(row=1, column=1).value == "Categoría"
    assert hoja.max_row == 2  # encabezado + 1 gasto
    assert hoja.cell(row=2, column=1).value == "Mantenimiento"
    assert hoja.cell(row=2, column=2).value == 1200.00


def test_export_budget_report_shows_planned_vs_actual(client):
    client.post(
        "/budgets",
        json={
            "categoria": "Mantenimiento", "periodicidad": "mensual",
            "periodo": "2026-09-01", "monto_planeado": 5000.00,
        },
    )
    client.post(
        "/expenses",
        json={
            "categoria": "Mantenimiento", "monto": 1200.00,
            "comprobante_url": "https://x/a.pdf", "fecha": "2026-09-05",
        },
    )

    wb = _leer_xlsx(client.get("/reports/budget/export", params={"periodo": "2026-09-01"}))
    hoja = wb["Presupuesto vs. real"]
    assert hoja.cell(row=2, column=1).value == "Mantenimiento"
    assert hoja.cell(row=2, column=3).value == 5000.00
    assert hoja.cell(row=2, column=4).value == 1200.00
    assert hoja.cell(row=2, column=5).value == 3800.00


def test_export_endpoints_default_when_no_data(client):
    wb = _leer_xlsx(client.get("/reports/expenses/export"))
    assert wb["Gastos"].max_row == 1  # solo el encabezado


@pytest.mark.asyncio
async def test_export_accounting_entries_includes_income_and_expenses(client, deposit_env):
    client.post("/properties", json={"identificador": "Casa 1"})
    client.post("/fees", json={"monto": 1500.00, "periodicidad": "mensual", "activa_desde": "2026-09-01"})
    client.post("/fees/generate-charges", params={"periodo": "2026-09-01"})

    async with deposit_env() as control_db:
        await dps.process_incoming_deposit(_deposito(), control_db)

    client.post(
        "/expenses",
        json={
            "categoria": "Mantenimiento", "monto": 800.00,
            "comprobante_url": "https://x/a.pdf", "fecha": "2026-09-10",
        },
    )

    wb = _leer_xlsx(client.get("/reports/accounting-entries/export"))
    hoja = wb["Movimientos"]
    assert hoja.cell(row=1, column=1).value == "Fecha"
    filas = [[hoja.cell(row=r, column=c).value for c in range(1, 7)] for r in range(2, hoja.max_row + 1)]
    tipos = {fila[1] for fila in filas}
    assert tipos == {"Ingreso", "Egreso"}
    ingreso = next(f for f in filas if f[1] == "Ingreso")
    assert ingreso[4] == 1500.00
    egreso = next(f for f in filas if f[1] == "Egreso")
    assert egreso[4] == 800.00
