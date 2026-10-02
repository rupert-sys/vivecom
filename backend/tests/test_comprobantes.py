"""
Comprobantes como archivo: el de un gasto (foto o PDF en lugar de un enlace de texto) y el que el
residente adjunta a su pago desde el banco, que tesorería revisa y acepta o rechaza.
"""

import uuid
from datetime import date, datetime, timedelta, timezone
from urllib.parse import urlsplit

import pytest

from app.api.deps import get_current_user
from app.core.config import settings
from app.main import app
from app.models.payment import EstadoPago, Payment
from app.schemas.auth import CurrentUser

JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 64
PDF = b"%PDF-1.7\n" + b"0" * 64


def _como(rol: str, user_id: str | None = None, property_id: str | None = None) -> str:
    tenant_id = app.dependency_overrides[get_current_user]().tenant_id
    uid = user_id or str(uuid.uuid4())

    def override():
        return CurrentUser(user_id=uid, tenant_id=tenant_id, schema_name="test", rol=rol, property_id=property_id)

    app.dependency_overrides[get_current_user] = override
    return uid


def _subir(client, kind="pago", contenido=JPEG, nombre="recibo.jpg") -> dict:
    respuesta = client.post("/files", files={"file": (nombre, contenido, "application/octet-stream")}, data={"kind": kind})
    assert respuesta.status_code == 201, respuesta.text
    return respuesta.json()


def _descargar(client, url: str):
    partes = urlsplit(url)
    return client.get(f"{partes.path}?{partes.query}")


def _gasto(**extra):
    base = {"categoria": "limpieza", "monto": 500, "fecha": "2026-09-01"}
    base.update(extra)
    return base


# ---------- Comprobante de un gasto ----------


def test_un_gasto_con_su_comprobante_subido_se_lee_con_un_enlace_que_lo_abre(client):
    _como("tesorero")  # F0-12: gastos son de tesorería, no del administrador (subir el comprobante incluido)
    archivo = _subir(client, kind="gasto", contenido=PDF, nombre="factura.pdf")
    creado = client.post("/expenses", json=_gasto(comprobante_archivo_id=archivo["id"], tipo_comprobante="factura"))

    assert creado.status_code == 201
    url = creado.json()["comprobante_url"]
    assert url.startswith(f"{settings.public_base_url}/files/{archivo['id']}/content?t=")

    assert _descargar(client, url).content == PDF
    listado = client.get("/expenses").json()
    assert _descargar(client, listado[0]["comprobante_url"]).status_code == 200
    detalle = client.get(f"/expenses/{creado.json()['id']}").json()
    assert _descargar(client, detalle["comprobante_url"]).status_code == 200


def test_un_gasto_sigue_aceptando_un_enlace_como_comprobante(client):
    _como("tesorero")
    creado = client.post("/expenses", json=_gasto(comprobante_url="https://ejemplo.com/factura.pdf"))
    assert creado.status_code == 201
    assert creado.json()["comprobante_url"] == "https://ejemplo.com/factura.pdf"  # un enlace externo queda tal cual


def test_un_gasto_sin_ningun_comprobante_se_rechaza(client):
    _como("tesorero")
    respuesta = client.post("/expenses", json=_gasto())
    assert respuesta.status_code == 422
    assert "comprobante" in respuesta.text.lower()


def test_el_comprobante_debe_ser_un_archivo_de_gasto_que_exista(client):
    _como("tesorero")
    assert client.post("/expenses", json=_gasto(comprobante_archivo_id=str(uuid.uuid4()))).status_code == 422
    de_pago = _subir(client, kind="pago")
    assert client.post("/expenses", json=_gasto(comprobante_archivo_id=de_pago["id"])).status_code == 422


def test_las_cotizaciones_tambien_pueden_ser_archivos(client):
    _como("tesorero")
    comprobante = _subir(client, kind="gasto")
    cotizacion = _subir(client, kind="gasto", contenido=PDF, nombre="cotizacion.pdf")
    creado = client.post(
        "/expenses",
        json=_gasto(
            comprobante_archivo_id=comprobante["id"],
            cotizaciones=[
                {"proveedor": "A", "monto": 100, "archivo_id": cotizacion["id"]},
                {"proveedor": "B", "monto": 120, "url": "https://ejemplo.com/b.pdf"},
            ],
        ),
    ).json()

    con_archivo, con_enlace = creado["cotizaciones"]
    assert _descargar(client, con_archivo["url"]).content == PDF
    assert con_enlace["url"] == "https://ejemplo.com/b.pdf"


def test_un_residente_abre_el_comprobante_de_un_gasto_pero_no_puede_registrar_uno(client):
    _como("tesorero")
    archivo = _subir(client, kind="gasto")
    client.post("/expenses", json=_gasto(comprobante_archivo_id=archivo["id"]))

    _como("residente", property_id=str(uuid.uuid4()))
    url = client.get("/expenses").json()[0]["comprobante_url"]
    assert _descargar(client, url).status_code == 200
    assert client.post("/expenses", json=_gasto(comprobante_url="https://x")).status_code == 403


def test_el_excel_de_gastos_no_lleva_la_referencia_interna_como_enlace(client):
    from io import BytesIO

    from openpyxl import load_workbook

    _como("tesorero")
    archivo = _subir(client, kind="gasto")
    client.post("/expenses", json=_gasto(comprobante_archivo_id=archivo["id"]))
    excel = load_workbook(BytesIO(client.get("/reports/expenses/export").content))
    celdas = [c.value for fila in excel.active.iter_rows() for c in fila]
    assert "Archivo adjunto en Vivecom" in celdas
    assert not any(isinstance(c, str) and c.startswith("/files/") for c in celdas)


# ---------- Comprobante de pago del residente ----------


def _vivienda_con_cargo(client, monto=750.0):
    prop = client.post("/properties", json={"identificador": "Casa 1"}).json()["id"]
    _como("tesorero")  # F0-12: cuotas y cargos son de tesorería, no del administrador
    client.post("/fees", json={"monto": monto, "periodicidad": "mensual", "activa_desde": "2026-01-01"})
    client.post("/fees/generate-charges", params={"periodo": "2026-09-01"})
    return prop


def _enviar(client, monto=750, **extra):
    archivo = _subir(client, kind="pago")
    cuerpo = {"monto": monto, "archivo_id": archivo["id"], **extra}
    return client.post("/payment-proofs", json=cuerpo)


def test_el_residente_envia_su_comprobante_y_queda_pendiente_sin_tocar_su_cuenta(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)

    enviado = _enviar(client, 750, fecha_pago="2026-09-04", nota="Transferencia BBVA")

    assert enviado.status_code == 201
    prueba = enviado.json()
    assert (prueba["estado"], prueba["monto"], prueba["nota"]) == ("pendiente", 750.0, "Transferencia BBVA")
    assert _descargar(client, prueba["archivo_url"]).content == JPEG
    _como("tesorero")
    assert client.get(f"/properties/{prop}/statement").json()["deuda_total"] == 750.0  # nada se concilió todavía


def test_solo_puede_adjuntar_un_archivo_propio_y_de_pago(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    ajeno = _subir(client, kind="pago")  # lo subió este residente...
    _como("residente", property_id=prop)  # ...y ahora es otra persona
    assert client.post("/payment-proofs", json={"monto": 750, "archivo_id": ajeno["id"]}).status_code == 422

    _como("tesorero")
    de_gasto = _subir(client, kind="gasto")
    _como("residente", property_id=prop)
    assert client.post("/payment-proofs", json={"monto": 750, "archivo_id": de_gasto["id"]}).status_code == 422
    assert client.post("/payment-proofs", json={"monto": 750, "archivo_id": str(uuid.uuid4())}).status_code == 422


def test_validaciones_del_comprobante(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    assert _enviar(client, monto=0).status_code == 422
    manana_lejano = (date.today() + timedelta(days=30)).isoformat()
    assert _enviar(client, fecha_pago=manana_lejano).status_code == 422

    _como("residente")  # sin vivienda
    assert client.post("/payment-proofs", json={"monto": 1, "archivo_id": str(uuid.uuid4())}).status_code == 400


def test_no_se_pueden_acumular_comprobantes_pendientes_sin_limite(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    for _ in range(20):
        assert _enviar(client).status_code == 201
    assert _enviar(client).status_code == 409


def test_cada_residente_ve_solo_los_de_su_vivienda_y_tesoreria_todos(client):
    una = _vivienda_con_cargo(client)
    _como("admin")  # dar de alta una vivienda sigue siendo del administrador
    otra = client.post("/properties", json={"identificador": "Casa 2"}).json()["id"]
    _como("residente", property_id=una)
    _enviar(client)
    _como("residente", property_id=otra)
    _enviar(client)
    assert len(client.get("/payment-proofs").json()) == 1

    _como("tesorero")
    assert len(client.get("/payment-proofs").json()) == 2
    assert len(client.get("/payment-proofs", params={"property_id": una}).json()) == 1
    assert client.get("/payment-proofs", params={"estado": "aceptado"}).json() == []


def test_tesoreria_acepta_el_comprobante_y_se_registra_el_pago_y_se_salda_el_cargo(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    prueba = _enviar(client, 750).json()

    tesorero = _como("tesorero")
    aceptado = client.post(f"/payment-proofs/{prueba['id']}/accept", json={})

    assert aceptado.status_code == 200
    resultado = aceptado.json()
    assert resultado["estado"] == "aceptado" and resultado["payment_id"] is not None
    estado = client.get(f"/properties/{prop}/statement").json()
    assert estado["deuda_total"] == 0.0 and estado["cargos"][0]["estado"] == "pagado"
    pago = next(p for p in client.get("/payments").json() if p["id"] == resultado["payment_id"])
    assert (pago["proveedor"], pago["registrado_por"], pago["estado"]) == ("transferencia", tesorero, "confirmado")

    # el residente ve su comprobante aceptado y puede descargar el recibo del pago
    _como("residente", property_id=prop)
    assert client.get("/payment-proofs").json()[0]["estado"] == "aceptado"
    assert client.get(f"/payments/{resultado['payment_id']}/receipt").status_code == 200


def test_tesoreria_puede_corregir_el_monto_que_de_verdad_llego(client):
    prop = _vivienda_con_cargo(client, monto=750)
    _como("residente", property_id=prop)
    prueba = _enviar(client, 1000).json()  # el residente se equivocó al capturar

    _como("tesorero")
    client.post(f"/payment-proofs/{prueba['id']}/accept", json={"monto": 750})

    assert client.get(f"/properties/{prop}/statement").json()["saldo_a_favor"] == 0.0


def test_un_comprobante_solo_se_revisa_una_vez(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    prueba = _enviar(client).json()
    _como("tesorero")
    assert client.post(f"/payment-proofs/{prueba['id']}/accept", json={}).status_code == 200
    assert client.post(f"/payment-proofs/{prueba['id']}/accept", json={}).status_code == 409
    assert client.post(f"/payment-proofs/{prueba['id']}/reject", json={"motivo": "x"}).status_code == 409
    assert len(client.get("/payments").json()) == 1  # el pago no se duplicó


def test_tesoreria_rechaza_con_un_motivo_y_el_residente_lo_ve(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    prueba = _enviar(client).json()

    _como("tesorero")
    assert client.post(f"/payment-proofs/{prueba['id']}/reject", json={"motivo": ""}).status_code == 422
    rechazado = client.post(f"/payment-proofs/{prueba['id']}/reject", json={"motivo": "Comprobante ilegible"})
    assert rechazado.status_code == 200
    assert client.get(f"/properties/{prop}/statement").json()["deuda_total"] == 750.0  # no se registró nada

    _como("residente", property_id=prop)
    visto = client.get("/payment-proofs").json()[0]
    assert (visto["estado"], visto["motivo_rechazo"]) == ("rechazado", "Comprobante ilegible")


def test_un_residente_no_puede_aceptar_ni_rechazar(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    prueba = _enviar(client).json()
    assert client.post(f"/payment-proofs/{prueba['id']}/accept", json={}).status_code == 403
    assert client.post(f"/payment-proofs/{prueba['id']}/reject", json={"motivo": "x"}).status_code == 403


def test_aceptar_o_rechazar_un_comprobante_inexistente_es_404(client):
    assert client.post(f"/payment-proofs/{uuid.uuid4()}/accept", json={}).status_code == 404
    assert client.post(f"/payment-proofs/{uuid.uuid4()}/reject", json={"motivo": "x"}).status_code == 404


@pytest.mark.asyncio
async def test_si_el_spei_ya_detecto_ese_pago_no_se_cuenta_dos_veces(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    prueba = _enviar(client, 750, fecha_pago=date.today().isoformat()).json()

    async with client.db_session_factory() as db:  # el SPEI ya detectó ese mismo depósito
        db.add(
            Payment(
                property_id=uuid.UUID(prop), monto=750, estado=EstadoPago.confirmado, referencia_recibida="0000001",
                clave_rastreo="SPEI-1", proveedor="stp",
                fecha_deteccion=datetime.now(timezone.utc).replace(tzinfo=None),
            )
        )
        await db.commit()

    _como("tesorero")
    duplicado = client.post(f"/payment-proofs/{prueba['id']}/accept", json={})
    assert duplicado.status_code == 409
    assert "SPEI ya detectó" in duplicado.json()["detail"]
    assert client.get("/payment-proofs").json()[0]["estado"] == "pendiente"

    # tesorería confirma que es otro pago (ej. de otro mes)
    forzado = client.post(f"/payment-proofs/{prueba['id']}/accept", json={"forzar": True})
    assert forzado.status_code == 200 and forzado.json()["estado"] == "aceptado"


@pytest.mark.asyncio
async def test_un_pago_capturado_a_mano_no_cuenta_como_ya_detectado(client):
    prop = _vivienda_con_cargo(client)
    _como("residente", property_id=prop)
    prueba = _enviar(client, 750).json()
    async with client.db_session_factory() as db:
        db.add(
            Payment(
                property_id=uuid.UUID(prop), monto=750, estado=EstadoPago.confirmado, referencia_recibida="0000001",
                clave_rastreo="MANUAL-x", proveedor="efectivo", registrado_por=uuid.uuid4(),
                fecha_deteccion=datetime.now(timezone.utc).replace(tzinfo=None),
            )
        )
        await db.commit()

    _como("tesorero")
    assert client.post(f"/payment-proofs/{prueba['id']}/accept", json={}).status_code == 200
