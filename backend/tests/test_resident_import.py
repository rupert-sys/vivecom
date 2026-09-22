"""
Importar la base de condóminos desde Excel: nombre, teléfono, correo, propietario o inquilino, número de casa.
Crea (o completa) viviendas y residentes y los liga — ver resident_import_service.py para las reglas. NO crea
cuentas de acceso: eso lo sigue haciendo el administrador aparte, desde Usuarios (test_users.py).
"""

import asyncio
import io
import uuid

from openpyxl import Workbook, load_workbook

from app.api.deps import get_current_user
from app.main import app
from app.models.resident import ResidentProperty, RolOcupacion
from app.schemas.auth import CurrentUser


def _tenant_id() -> str:
    return app.dependency_overrides[get_current_user]().tenant_id


def _como(rol: str) -> None:
    tenant_id = _tenant_id()

    def override():
        return CurrentUser(user_id=str(uuid.uuid4()), tenant_id=tenant_id, schema_name="test", rol=rol, property_id=None)

    app.dependency_overrides[get_current_user] = override


ENCABEZADOS = ["Nombre", "Teléfono", "Correo", "Propietario o inquilino", "Número de casa"]


def _excel(filas: list[list], encabezados: list[str] | None = ENCABEZADOS) -> bytes:
    wb = Workbook()
    ws = wb.active
    if encabezados is not None:
        ws.append(encabezados)
    for fila in filas:
        ws.append(fila)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _importar(client, filas, encabezados=ENCABEZADOS, nombre_archivo="condominos.xlsx"):
    contenido = _excel(filas, encabezados)
    return client.post(
        "/residents/import", files={"archivo": (nombre_archivo, contenido, "application/octet-stream")}
    )


def test_importa_propietarios_e_inquilinos_y_crea_viviendas_y_residentes(client):
    r = _importar(
        client,
        [
            ["Mariana Ortega", "5555550101", "mariana@condo.mx", "propietario", "Casa 1"],
            ["Luis Peña", "5555550102", "", "inquilino", "Casa 2"],  # el correo es opcional
        ],
    )
    assert r.status_code == 200
    body = r.json()
    assert (body["total_filas"], body["viviendas_creadas"], body["residentes_creados"]) == (2, 2, 2)
    assert (body["vinculos_creados"], body["filas_con_error"]) == (2, [])

    _como("admin")
    viviendas = {p["identificador"]: p["id"] for p in client.get("/properties").json()}
    residentes = {r["nombre"]: r for r in client.get("/residents").json()}
    assert set(viviendas) == {"Casa 1", "Casa 2"}
    assert residentes["Mariana Ortega"]["email"] == "mariana@condo.mx"

    ligados = client.get(f"/properties/{viviendas['Casa 2']}/residents").json()
    assert [p["nombre"] for p in ligados] == ["Luis Peña"]


def test_reconoce_encabezados_con_acentos_mayusculas_y_variantes(client):
    r = _importar(
        client,
        [["Ana Duarte", "5555550103", "ana@condo.mx", "Dueña", "Casa 3"]],
        encabezados=["  NOMBRE  ", "Tel", "E-mail", "Tipo", "No. de casa"],
    )
    assert r.status_code == 200
    assert r.json()["residentes_creados"] == 1


def test_reintentar_el_mismo_archivo_no_duplica_nada(client):
    filas = [["Mariana Ortega", "5555550101", "mariana@condo.mx", "propietario", "Casa 1"]]
    assert _importar(client, filas).json()["viviendas_creadas"] == 1

    segunda = _importar(client, filas).json()
    assert (segunda["viviendas_creadas"], segunda["residentes_creados"], segunda["vinculos_creados"]) == (0, 0, 0)
    _como("admin")
    assert len(client.get("/properties").json()) == 1
    assert len(client.get("/residents").json()) == 1


def test_el_mismo_correo_en_dos_filas_actualiza_al_residente_en_vez_de_duplicarlo(client):
    r = _importar(
        client,
        [
            ["Mariana Ortega", "5555550101", "mariana@condo.mx", "propietario", "Casa 1"],
            ["Mariana Ortega Salazar", "5555550199", "mariana@condo.mx", "propietario", "Casa 2"],  # se mudó / dato más nuevo
        ],
    )
    body = r.json()
    assert (body["residentes_creados"], body["residentes_actualizados"]) == (1, 1)
    _como("admin")
    residentes = client.get("/residents").json()
    assert len(residentes) == 1
    assert (residentes[0]["nombre"], residentes[0]["telefono"]) == ("Mariana Ortega Salazar", "5555550199")


def test_si_cambia_el_rol_en_una_reimportacion_se_actualiza_el_vinculo(client):
    filas = [["Mariana Ortega", "5555550101", "mariana@condo.mx", "propietario", "Casa 1"]]
    _importar(client, filas)

    otras = [["Mariana Ortega", "5555550101", "mariana@condo.mx", "inquilino", "Casa 1"]]
    r = _importar(client, otras)
    assert (r.json()["vinculos_creados"], r.json()["vinculos_actualizados"]) == (0, 1)

    _como("admin")
    vivienda_id = client.get("/properties").json()[0]["id"]
    residente_id = client.get("/residents").json()[0]["id"]
    async def _rol():
        async with client.db_session_factory() as db:
            link = await db.get(ResidentProperty, {"resident_id": uuid.UUID(residente_id), "property_id": uuid.UUID(vivienda_id)})
            return link.rol

    assert asyncio.run(_rol()) == RolOcupacion.inquilino


def test_reintentar_sin_correo_tampoco_duplica_al_residente(client):
    """Es el caso más común: el Excel del condominio rara vez trae el correo de todos."""
    filas = [["Luis Peña", "5555550102", "", "inquilino", "Casa 2"]]
    assert _importar(client, filas).json()["residentes_creados"] == 1

    segunda = _importar(client, filas).json()
    assert (segunda["residentes_creados"], segunda["residentes_actualizados"]) == (0, 0)
    _como("admin")
    assert len([r for r in client.get("/residents").json() if r["nombre"] == "Luis Peña"]) == 1


def test_dos_residentes_sin_correo_con_distinto_nombre_no_se_confunden(client):
    r = _importar(
        client,
        [
            ["Luis Peña", "5555550102", "", "inquilino", "Casa 2"],
            ["Carmen Beltrán", "5555550106", "", "inquilino", "Casa 6"],
        ],
    )
    assert r.json()["residentes_creados"] == 2


def test_filas_invalidas_se_reportan_y_no_detienen_las_demas(client):
    r = _importar(
        client,
        [
            ["", "5555550101", "", "propietario", "Casa 1"],  # sin nombre
            ["Luis Peña", "", "", "inquilino", "Casa 2"],  # sin teléfono
            ["Ana Duarte", "5555550103", "", "vecino", "Casa 3"],  # rol inválido
            ["Jorge Ramírez", "5555550104", "no-es-correo", "propietario", "Casa 4"],  # correo inválido
            ["Sofía Navarro", "5555550105", "", "propietario", ""],  # sin vivienda
            ["Carmen Beltrán", "5555550106", "carmen@condo.mx", "propietario", "Casa 6"],  # esta sí es válida
        ],
    )
    body = r.json()
    assert body["total_filas"] == 6
    assert (body["viviendas_creadas"], body["residentes_creados"]) == (1, 1)
    filas_con_error = {e["fila"] for e in body["filas_con_error"]}
    assert filas_con_error == {2, 3, 4, 5, 6}
    motivos = {e["fila"]: e["motivo"] for e in body["filas_con_error"]}
    assert "nombre" in motivos[2].lower()
    assert "teléfono" in motivos[3].lower() or "telefono" in motivos[3].lower()
    assert "vecino" in motivos[4]
    assert "no-es-correo" in motivos[5]
    assert "casa" in motivos[6].lower()


def test_filas_en_blanco_al_final_no_cuentan_ni_son_error(client):
    r = _importar(
        client,
        [
            ["Mariana Ortega", "5555550101", "mariana@condo.mx", "propietario", "Casa 1"],
            [None, None, None, None, None],
            ["", "", "", "", ""],
        ],
    )
    body = r.json()
    assert (body["total_filas"], body["filas_con_error"]) == (1, [])


def test_encabezados_que_no_se_reconocen_se_rechazan_con_un_solo_mensaje_claro(client):
    r = _importar(client, [["x", "y"]], encabezados=["Columna A", "Columna B"])
    assert r.status_code == 422
    assert "plantilla" in r.json()["detail"].lower()


def test_un_excel_sin_ninguna_fila_de_datos_no_es_error_solo_no_importa_nada(client):
    r = _importar(client, [])  # encabezados sí, ninguna fila de condóminos abajo
    assert r.status_code == 200
    assert r.json()["total_filas"] == 0


def test_un_archivo_vacio_o_que_no_es_excel_se_rechaza(client):
    r = _importar(client, [], encabezados=None)  # ni encabezados: la hoja está completamente vacía
    assert r.status_code == 422 and "vacío" in r.json()["detail"].lower()

    r2 = client.post("/residents/import", files={"archivo": ("condominos.xlsx", b"esto no es un excel", "application/octet-stream")})
    assert r2.status_code == 422

    r3 = client.post("/residents/import", files={"archivo": ("condominos.csv", b"a,b", "text/csv")})
    assert r3.status_code == 422 and "excel" in r3.json()["detail"].lower()


def test_solo_el_admin_importa_y_descarga_la_plantilla(client):
    for rol in ("tesorero", "comite_aprobador", "residente", "guardia"):
        _como(rol)
        assert _importar(client, [["a", "b", "c", "propietario", "Casa 1"]]).status_code == 403
        assert client.get("/residents/import-template").status_code == 403


def test_la_plantilla_se_puede_reimportar_tal_cual(client):
    plantilla = client.get("/residents/import-template")
    assert plantilla.status_code == 200
    assert plantilla.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    wb = load_workbook(io.BytesIO(plantilla.content))
    ws = wb.active
    filas = list(ws.iter_rows(values_only=True))[1:]  # sin el encabezado, ya viene con una fila de ejemplo
    r = client.post(
        "/residents/import", files={"archivo": ("plantilla-condominos.xlsx", plantilla.content, "application/octet-stream")}
    )
    assert r.status_code == 200
    assert r.json()["filas_con_error"] == []
    assert r.json()["total_filas"] == len(filas) == 1
