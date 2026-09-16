"""
Exportación a Excel (F1-17, HU-A11): estado de cuenta, gastos y presupuesto
vs. real, filtrables por vivienda/categoría/rango de fechas.

Se genera de forma síncrona dentro del propio request, no vía Celery: el
plan la describe como "asíncrona", pero este proyecto no tiene todavía
infraestructura de almacenamiento de archivos ni de seguimiento de jobs (ver
comprobante_url en expense.py — el mismo hueco), y el volumen de datos de un
condominio no justifica construir esa cola de "generar y avisar cuando esté
listo" para lo que en la práctica es un reporte pequeño. Se sirve como
descarga directa, igual que el recibo de F1-11.
"""

import io
import uuid
from datetime import date

from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.expense import Expense
from app.models.payment import EstadoPago, Payment
from app.models.property import Property
from app.services.budget_service import get_budget_vs_actual
from app.services.statement_service import get_account_statement


def _workbook_to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def build_account_statements_excel(db: AsyncSession, property_id: uuid.UUID | None) -> bytes:
    if property_id is not None:
        propiedad = await db.get(Property, property_id)
        propiedades = [propiedad] if propiedad is not None else []
    else:
        propiedades = (await db.execute(select(Property).order_by(Property.identificador))).scalars().all()

    wb = Workbook()
    hoja_cargos = wb.active
    hoja_cargos.title = "Cargos"
    hoja_cargos.append(["Vivienda", "Periodo", "Monto base", "Recargo", "Total", "Estado"])

    hoja_pagos = wb.create_sheet("Pagos")
    hoja_pagos.append(["Vivienda", "Fecha", "Monto", "Estado", "Clave de rastreo"])

    for propiedad in propiedades:
        estado = await get_account_statement(db, propiedad.id)
        for cargo in estado.cargos:
            total = float(cargo.monto_base) + float(cargo.recargo_aplicado)
            hoja_cargos.append(
                [propiedad.identificador, cargo.periodo, float(cargo.monto_base), float(cargo.recargo_aplicado),
                 total, cargo.estado.value]
            )
        for pago in estado.pagos:
            hoja_pagos.append(
                [propiedad.identificador, pago.fecha_deteccion, float(pago.monto), pago.estado.value,
                 pago.clave_rastreo]
            )

    return _workbook_to_bytes(wb)


async def build_expenses_excel(
    db: AsyncSession, desde: date | None, hasta: date | None, categoria: str | None
) -> bytes:
    query = select(Expense).order_by(Expense.fecha)
    if desde is not None:
        query = query.where(Expense.fecha >= desde)
    if hasta is not None:
        query = query.where(Expense.fecha <= hasta)
    if categoria is not None:
        query = query.where(Expense.categoria == categoria)
    gastos = (await db.execute(query)).scalars().all()

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Gastos"
    hoja.append(["Categoría", "Monto", "Fecha", "Comprobante"])
    for gasto in gastos:
        hoja.append([gasto.categoria, float(gasto.monto), gasto.fecha, gasto.comprobante_url])

    return _workbook_to_bytes(wb)


async def build_accounting_entries_excel(db: AsyncSession, desde: date | None, hasta: date | None) -> bytes:
    """
    F3-01: "solo exportación de pólizas para que el contador la use en su
    propio sistema" (alcance §3.1) — explícitamente NO contabilidad de
    doble entrada. Cada movimiento (ingreso o egreso) sale como una fila
    plana; el contador la importa como auxiliar y completa la partida
    contraria (banco/caja) en su propio software.

    AVISO: no está verificado contra el formato de importación real de
    CONTPAQi ni Aspel (requeriría una licencia/cuenta de alguno de los dos
    para confirmar su esquema exacto de columnas) — mismo espíritu que el
    aviso en payment_providers/stp.py. Se entrega un Excel con columnas
    genéricas (Fecha/Tipo/Cuenta/Concepto/Monto/Referencia) que cualquier
    asistente de importación puede re-mapear, en vez de arriesgar un
    formato propietario específico sin poder confirmarlo. "Cuenta" usa la
    categoría de Expense o una cuenta de ingresos fija — no hay catálogo de
    cuentas contables real en este modelo de datos; el contador de cada
    condominio deberá mapear estas columnas a su propio catálogo.
    """
    ingresos_query = select(Payment, Property.identificador).join(
        Property, Property.id == Payment.property_id
    ).where(Payment.estado == EstadoPago.confirmado)
    if desde is not None:
        ingresos_query = ingresos_query.where(Payment.fecha_deteccion >= desde)
    if hasta is not None:
        ingresos_query = ingresos_query.where(Payment.fecha_deteccion <= hasta)
    ingresos = (await db.execute(ingresos_query.order_by(Payment.fecha_deteccion))).all()

    egresos_query = select(Expense).order_by(Expense.fecha)
    if desde is not None:
        egresos_query = egresos_query.where(Expense.fecha >= desde)
    if hasta is not None:
        egresos_query = egresos_query.where(Expense.fecha <= hasta)
    egresos = (await db.execute(egresos_query)).scalars().all()

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Movimientos"
    hoja.append(["Fecha", "Tipo", "Cuenta", "Concepto", "Monto", "Referencia"])

    for pago, identificador in ingresos:
        hoja.append(
            [pago.fecha_deteccion, "Ingreso", "Ingresos por cuotas de mantenimiento",
             f"Cuota — {identificador}", float(pago.monto), pago.clave_rastreo]
        )
    for gasto in egresos:
        hoja.append([gasto.fecha, "Egreso", gasto.categoria, gasto.categoria, float(gasto.monto), str(gasto.id)])

    return _workbook_to_bytes(wb)


async def build_budget_report_excel(db: AsyncSession, periodo: date) -> bytes:
    comparaciones = await get_budget_vs_actual(db, periodo)

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Presupuesto vs. real"
    hoja.append(["Categoría", "Periodicidad", "Planeado", "Real", "Diferencia"])
    for comparacion in comparaciones:
        hoja.append(
            [comparacion.categoria, comparacion.periodicidad.value, comparacion.monto_planeado,
             comparacion.monto_real, comparacion.monto_planeado - comparacion.monto_real]
        )

    return _workbook_to_bytes(wb)
