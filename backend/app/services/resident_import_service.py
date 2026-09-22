"""
Importar la base de condóminos desde un Excel: nombre, teléfono, correo, propietario o inquilino, y el número de
casa. Con eso crea (o completa) las Property y los Resident y los liga (ResidentProperty) — ver esos modelos para
el porqué de esa forma. NO crea cuentas de acceso (UserAccount): dar de alta quién puede entrar al panel o a la
app es una decisión aparte, deliberada, que el administrador sigue haciendo desde Usuarios (con su contraseña).

No se detiene en la primera fila mal capturada: procesa las que sí puede y regresa cuáles fallaron y por qué, para
que el administrador corrija solo esas en el Excel y lo vuelva a subir — reintentable: una vivienda o un residente
que ya existían no se duplican, se completan.

Un correo solo identifica a la MISMA persona si además el nombre coincide (ver el bloque de identidad más abajo):
sin esta exigencia, varias filas con un correo repetido por accidente (p. ej. el de ejemplo de la plantilla, sin
corregir) se fusionaban en un solo residente inventado, "dueño" de todas esas viviendas — bug real encontrado en
producción. Un choque de correo con otro nombre se reporta como fila con error, nunca se fusiona en automático.
"""

import io
import re
import unicodedata

from openpyxl import Workbook, load_workbook
from pydantic import EmailStr, TypeAdapter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.property import Property
from app.models.resident import Resident, ResidentProperty, RolOcupacion
from app.schemas.resident_import import ResidentImportError, ResidentImportResult
from app.services.payment_reference import generate_payment_reference

# Varios nombres de columna aceptados (el administrador exporta de donde tenga la lista, no siempre igual):
# la clave es el campo interno, cada valor es una forma normalizada (ver _normalizar) del encabezado.
_ENCABEZADOS = {
    "nombre": {"nombre", "nombre completo"},
    "telefono": {"telefono", "tel", "celular", "whatsapp"},
    "correo": {"correo", "correo electronico", "email", "e mail", "mail"},
    "rol": {"propietario o inquilino", "propietario inquilino", "rol", "tipo", "ocupacion"},
    "vivienda": {"numero de casa", "no de casa", "no de casa.", "casa", "vivienda", "identificador", "lote"},
}
# Claves ya normalizadas (ver _normalizar: sin acentos, así que "dueño"/"dueña" llegan aquí como "dueno"/"duena").
_ROLES = {
    "propietario": RolOcupacion.propietario, "dueno": RolOcupacion.propietario, "duena": RolOcupacion.propietario,
    "inquilino": RolOcupacion.inquilino, "arrendatario": RolOcupacion.inquilino, "renta": RolOcupacion.inquilino,
    "rentando": RolOcupacion.inquilino,
}
_correo_valido = TypeAdapter(EmailStr)


def _normalizar(texto: str) -> str:
    """minúsculas, sin acentos, separadores (-_/.) y espacios de más colapsados: "E-mail", "Propietario/Inquilino"
    y "No. de casa" casan igual que "e mail", "propietario inquilino" y "no de casa"."""
    sin_acentos = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    sin_separadores = re.sub(r"[-_/.]+", " ", sin_acentos.strip().lower())
    return re.sub(r"\s+", " ", sin_separadores).strip()


def _mapear_columnas(encabezados: list[str | None]) -> dict[str, int]:
    normalizados = [_normalizar(h) if h else "" for h in encabezados]
    columnas: dict[str, int] = {}
    for campo, variantes in _ENCABEZADOS.items():
        for indice, encabezado in enumerate(normalizados):
            if encabezado in variantes:
                columnas[campo] = indice
                break
    faltantes = _ENCABEZADOS.keys() - columnas.keys()
    if faltantes:
        nombres = ", ".join(sorted(faltantes))
        raise ValueError(
            f"El Excel no tiene estas columnas (o no se reconoce su encabezado): {nombres}. "
            "Descarga la plantilla desde 'Importar Excel' para ver los nombres exactos."
        )
    return columnas


def _texto(valor: object) -> str:
    return str(valor).strip() if valor is not None else ""


async def importar_condominos(db: AsyncSession, contenido: bytes) -> ResidentImportResult:
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro.active
        filas = list(hoja.iter_rows(values_only=True))
    except Exception as exc:  # noqa: BLE001 — cualquier archivo corrupto o que no sea un .xlsx real cae aquí
        raise ValueError("No se pudo leer el archivo: verifica que sea un Excel (.xlsx) válido.") from exc
    if not filas:
        raise ValueError("El Excel está vacío.")

    columnas = _mapear_columnas(list(filas[0]))

    # Viviendas y residentes que ya existen, para no duplicarlos ni volver a consultar la BD en cada fila.
    viviendas = {p.identificador: p for p in (await db.execute(select(Property))).scalars().all()}
    todos_los_residentes = (await db.execute(select(Resident))).scalars().all()
    residentes_por_correo = {r.email.lower(): r for r in todos_los_residentes if r.email}
    # Respaldo para cuando la fila no trae correo (el caso más común: el Excel del condominio rara vez lo tiene
    # completo): sin esto, reimportar el mismo archivo creaba un residente nuevo cada vez en vez de reconocerlo.
    residentes_por_nombre_y_telefono = {(_normalizar(r.nombre), r.telefono.strip()): r for r in todos_los_residentes}
    vinculos = {
        (v.resident_id, v.property_id): v for v in (await db.execute(select(ResidentProperty))).scalars().all()
    }

    resultado = ResidentImportResult(
        total_filas=0, viviendas_creadas=0, residentes_creados=0, residentes_actualizados=0,
        vinculos_creados=0, vinculos_actualizados=0, filas_con_error=[],
    )

    for numero_de_fila, fila in enumerate(filas[1:], start=2):
        if fila is None or all(c is None or str(c).strip() == "" for c in fila):
            continue  # fila en blanco: ni cuenta ni es error, es normal al final de un Excel
        resultado.total_filas += 1

        def valor(campo: str) -> str:
            return _texto(fila[columnas[campo]]) if columnas[campo] < len(fila) else ""

        nombre, telefono, correo_txt, rol_txt, identificador = (
            valor("nombre"), valor("telefono"), valor("correo"), valor("rol"), valor("vivienda")
        )

        if not nombre:
            resultado.filas_con_error.append(ResidentImportError(fila=numero_de_fila, motivo="Falta el nombre."))
            continue
        if not identificador:
            resultado.filas_con_error.append(ResidentImportError(fila=numero_de_fila, motivo="Falta el número de casa."))
            continue
        if not telefono:
            resultado.filas_con_error.append(ResidentImportError(fila=numero_de_fila, motivo="Falta el teléfono."))
            continue
        rol = _ROLES.get(_normalizar(rol_txt))
        if rol is None:
            resultado.filas_con_error.append(
                ResidentImportError(fila=numero_de_fila, motivo=f'"{rol_txt}" no es "propietario" ni "inquilino".')
            )
            continue
        correo = None
        if correo_txt:
            try:
                correo = _correo_valido.validate_python(correo_txt).lower()
            except ValueError:
                resultado.filas_con_error.append(
                    ResidentImportError(fila=numero_de_fila, motivo=f'"{correo_txt}" no es un correo válido.')
                )
                continue

        # La identidad se resuelve ANTES de crear la vivienda: si el correo ya es de alguien más, la fila es un
        # error y no debe dejar una vivienda huérfana a medias (sin residente) que el admin tenga que notar aparte.
        clave_nombre_telefono = (_normalizar(nombre), telefono)
        residente = None
        if correo:
            existente = residentes_por_correo.get(correo)
            if existente is not None:
                if _normalizar(existente.nombre) == _normalizar(nombre):
                    residente = existente
                else:
                    # Bug real (F1-4x): varias filas con el correo de ejemplo de la plantilla, sin corregirlo, se
                    # fusionaban en un solo residente que terminaba "ligado" a decenas de viviendas ajenas. Un
                    # correo solo identifica a la MISMA persona si además el nombre coincide; si no, es un choque
                    # que hay que corregir a mano, no algo que se pueda adivinar y fusionar en automático.
                    resultado.filas_con_error.append(ResidentImportError(
                        fila=numero_de_fila,
                        motivo=(
                            f'El correo "{correo}" ya está registrado a nombre de "{existente.nombre}". Si es la '
                            "misma persona, usa exactamente ese nombre; si es otra, corrige el correo."
                        ),
                    ))
                    continue
        else:
            residente = residentes_por_nombre_y_telefono.get(clave_nombre_telefono)

        vivienda = viviendas.get(identificador)
        if vivienda is None:
            vivienda = Property(identificador=identificador, referencia_pago=await generate_payment_reference(identificador, db))
            db.add(vivienda)
            # A diferencia de un solo POST /properties (que hace commit antes de leer el id), aquí se siguen
            # creando filas en el mismo lote: hace falta el id YA para poder ligar la vivienda más abajo.
            await db.flush()
            viviendas[identificador] = vivienda
            resultado.viviendas_creadas += 1

        if residente is None:
            residente = Resident(nombre=nombre, telefono=telefono, email=correo)
            db.add(residente)
            await db.flush()
            if correo:
                residentes_por_correo[correo] = residente
            residentes_por_nombre_y_telefono[clave_nombre_telefono] = residente
            resultado.residentes_creados += 1
        elif residente.telefono != telefono:
            # El nombre ya coincidía (si venía por correo, se exigió arriba; si venía por nombre+teléfono, es la
            # propia clave de búsqueda) — lo único que puede traer de más nuevo esta fila es el teléfono.
            residente.telefono = telefono
            residentes_por_nombre_y_telefono[clave_nombre_telefono] = residente
            resultado.residentes_actualizados += 1

        clave_vinculo = (residente.id, vivienda.id)
        vinculo = vinculos.get(clave_vinculo)
        if vinculo is None:
            vinculos[clave_vinculo] = ResidentProperty(resident_id=residente.id, property_id=vivienda.id, rol=rol)
            db.add(vinculos[clave_vinculo])
            resultado.vinculos_creados += 1
        elif vinculo.rol != rol:
            vinculo.rol = rol
            resultado.vinculos_actualizados += 1

    await db.commit()
    return resultado


def construir_plantilla() -> bytes:
    """El Excel de ejemplo que se ofrece para descargar antes de importar: mismos encabezados, una fila de muestra."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Condóminos"
    ws.append(["Nombre", "Teléfono", "Correo", "Propietario o inquilino", "Número de casa"])
    ws.append(["Mariana Ortega Salazar", "5555550101", "mariana@example.com", "propietario", "Casa 1"])
    for columna, ancho in zip("ABCDE", (28, 16, 28, 22, 16), strict=True):
        ws.column_dimensions[columna].width = ancho
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
